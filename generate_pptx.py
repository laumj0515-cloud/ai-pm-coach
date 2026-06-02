from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ── Colors ──
BG_DARK  = RGBColor(0x0B, 0x11, 0x20)
BG_CARD  = RGBColor(0x11, 0x18, 0x27)
BRAND    = RGBColor(0x63, 0x66, 0xF1)
BRAND2   = RGBColor(0x81, 0x8C, 0xF8)
WHITE    = RGBColor(0xE2, 0xE8, 0xF0)
GRAY     = RGBColor(0x94, 0xA3, 0xB8)
GRAY2    = RGBColor(0x64, 0x74, 0x8B)
RED      = RGBColor(0xF8, 0x71, 0x71)
AMBER    = RGBColor(0xFB, 0xBF, 0x24)
GREEN    = RGBColor(0x34, 0xD3, 0x99)
BORDER   = RGBColor(0x1E, 0x2D, 0x3D)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

W = prs.slide_width
H = prs.slide_height

# ── Helper functions ──
def add_bg(slide, color=BG_DARK):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_rect(slide, left, top, width, height, fill_color=None, border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    return shape

def add_rounded_rect(slide, left, top, width, height, fill_color=None, border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    return shape

def add_text_box(slide, left, top, width, height, text, font_size=14, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name='Noto Sans SC'):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox

def add_section_label(slide, text, left=Inches(1), top=Inches(0.6)):
    add_text_box(slide, left, top, Inches(4), Inches(0.4), text, font_size=11, color=BRAND2, bold=True)

def add_title(slide, text, left=Inches(1), top=Inches(1.1)):
    add_text_box(slide, left, top, Inches(10), Inches(0.7), text, font_size=28, color=WHITE, bold=True)

def add_card(slide, left, top, width, height, icon, title, desc):
    card = add_rounded_rect(slide, left, top, width, height, fill_color=BG_CARD, border_color=BORDER)
    add_text_box(slide, left + Inches(0.3), top + Inches(0.3), width - Inches(0.6), Inches(0.5), icon, font_size=22)
    add_text_box(slide, left + Inches(0.3), top + Inches(0.9), width - Inches(0.6), Inches(0.4), title, font_size=14, color=WHITE, bold=True)
    add_text_box(slide, left + Inches(0.3), top + Inches(1.4), width - Inches(0.6), height - Inches(1.7), desc, font_size=11, color=GRAY2)

def add_dot_line(slide, left, top, width, color=BRAND2):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(2))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()

# ═══════════════ SLIDE 1: Title ═══════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_bg(slide)

# Big brand accent bar
add_rect(slide, Inches(0), Inches(0), Inches(0.08), H, fill_color=BRAND)

add_section_label(slide, 'AI PM · 面试训练系统')
add_text_box(slide, Inches(1), Inches(1.8), Inches(10), Inches(1.2),
    'AI PM Coach', font_size=64, color=WHITE, bold=True)
add_text_box(slide, Inches(1), Inches(2.9), Inches(10), Inches(0.8),
    'AI 面试训练系统 — 产品说明书', font_size=22, color=GRAY)
add_text_box(slide, Inches(1), Inches(4.2), Inches(10), Inches(0.5),
    '专为 AI 产品经理岗位设计的个性化模拟面试工具', font_size=16, color=GRAY2)
add_text_box(slide, Inches(1), Inches(5.8), Inches(6), Inches(0.4),
    '作者：刘美姣  |  v1.0  |  React · DeepSeek API · Web Speech', font_size=12, color=GRAY2)

# ═══════════════ SLIDE 2: 项目背景 ═══════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_section_label(slide, '01 · 项目背景')
add_title(slide, '为什么需要这个工具')
add_text_box(slide, Inches(1), Inches(1.9), Inches(10), Inches(0.5),
    '刘美姣，2年产品经验，家居行业转AI产品经理。一年Gap期独立开发Tempo和Mibo。求职中发现四个核心痛点：',
    font_size=13, color=GRAY)

cards_data = [
    ('🎯', '岗位名不副实', 'JD写"应用层AI PM"，面试问的全是模型层问题。需要练如何识别并回应——"这不是我该答的方向"。'),
    ('🔄', '表达能力弱环', '思维强但表达跳跃。面试时把内心逻辑密度变成信息过载，说很多但没说到点上。需训练结构化输出。'),
    ('📊', '数据支撑不够', 'Tempo和Mibo有真实数据和产品沉淀，但面试时无法快速调用。需反复训练"用数据说话"的肌肉记忆。'),
    ('📝', '无法模拟追问', '跟自己练习看不出问题。需要一个严格、会打断、会追问"不够具体"的真实面试官来对练。'),
]
card_w = Inches(2.6)
card_h = Inches(3.8)
start_x = Inches(1)
gap = Inches(0.25)
for i, (icon, title, desc) in enumerate(cards_data):
    add_card(slide, start_x + i * (card_w + gap), Inches(2.6), card_w, card_h, icon, title, desc)

# ═══════════════ SLIDE 3: 解决方案 ═══════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_section_label(slide, '02 · 解决方案')
add_title(slide, '一个记住你全部底牌的 AI 面试教练')

points = [
    ('不是问题库，而是对话引擎', '它需要听懂你说了什么，判断你是否真的回答到了问题——这是最难的地方，也是最重要的。'),
    ('不是鼓励型，而是严苛型', '因为你对敷衍极敏感，假装满意的反馈对你没用。它会打断你、追问你、说"不够具体"。'),
    ('不止练表达，还练边界感', '遇到模型训推、embedding这类不属于应用层PM的问题，它会告诉你"这不是你该答的方向"。'),
]
for i, (title, desc) in enumerate(points):
    y = Inches(2.5) + i * Inches(1.5)
    add_dot_line(slide, Inches(1), y, Inches(0.6), BRAND)
    add_text_box(slide, Inches(1), y + Inches(0.15), Inches(10), Inches(0.4), title, font_size=18, color=WHITE, bold=True)
    add_text_box(slide, Inches(1), y + Inches(0.55), Inches(9), Inches(0.5), desc, font_size=13, color=GRAY)

# ═══════════════ SLIDE 4: 五步闭环 ═══════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_section_label(slide, '03 · 核心功能')
add_title(slide, '五步闭环训练流程')

steps = [
    ('01', '选择模式', '创始人/HR/\n技术PM/薪资'),
    ('02', 'AI 提问', '基于你的背景\n深度追问'),
    ('03', '你回答', '文字/语音/\n长按录音'),
    ('04', '即时点评', '✅好的 ❌改善\n💡示范'),
    ('05', '总结复盘', '总分+逐题分析\n+薄弱追踪'),
]
step_w = Inches(2)
step_h = Inches(2.8)
total_w = len(steps) * step_w + (len(steps)-1) * Inches(0.3)
start_x = int((W - total_w) / 2)

for i, (num, label, sub) in enumerate(steps):
    x = start_x + i * (step_w + Inches(0.3))
    card = add_rounded_rect(slide, x, Inches(3), step_w, step_h, fill_color=BG_CARD, border_color=BORDER)
    add_text_box(slide, x + Inches(0.2), Inches(3.2), step_w - Inches(0.4), Inches(0.7), num, font_size=36, color=BRAND, bold=True)
    add_text_box(slide, x + Inches(0.2), Inches(4.0), step_w - Inches(0.4), Inches(0.4), label, font_size=16, color=WHITE, bold=True)
    add_text_box(slide, x + Inches(0.2), Inches(4.5), step_w - Inches(0.4), Inches(1.0), sub, font_size=12, color=GRAY2)
    # Arrow
    if i < len(steps) - 1:
        add_text_box(slide, x + step_w + Inches(0.02), Inches(4.1), Inches(0.3), Inches(0.4), '→', font_size=20, color=GRAY2)

# ═══════════════ SLIDE 5: 四种模式 ═══════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_section_label(slide, '04 · 面试模式')
add_title(slide, '四种模式，覆盖全部面试场景')

modes = [
    ('🚀', '初创创始人', 'CEO/创始人视角\n追问你能不能独立扛事\n有没有0到1的能力'),
    ('💼', 'HR 面试官', 'HR视角\n追问gap、薪资预期\n稳定性与团队匹配度'),
    ('⚙️', '技术PM负责人', '技术产品视角\n深挖架构逻辑与AI理解\n不越界到模型层'),
    ('💰', '薪资谈判', 'HR谈薪资模拟\nAI有隐藏预算\n你报价，看结果'),
]
for i, (icon, title, desc) in enumerate(modes):
    x = Inches(1) + i * Inches(2.9)
    add_card(slide, x, Inches(2.5), Inches(2.6), Inches(3.6), icon, title, desc)

# ═══════════════ SLIDE 6: 实时反馈 ═══════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_section_label(slide, '05 · 核心亮点')
add_title(slide, '每题即时反馈，而非一次性复盘')

# Feedback demo cards
feedback_cards = [
    ('✅ 做得好的', GREEN, 'BG_CARD fill + green border', '你用了Tempo的82%完成率来支撑论点，数据引用准确。开场先说结论再展开，逻辑结构清晰。'),
    ('❌ 需要改善', RED, 'BG_CARD fill + red border', '"赋能"这个词太虚，面试官听到没感觉。换成"通过五象限路由把完成率从45%拉到82%"。另外第二次追问时你绕开了问题——面试官问你为什么不是3象限，你说了一堆但没直接回答。'),
    ('💡 更好的答法', BRAND, 'BG_CARD fill + brand border', '"我选5象限而不是3象限，是因为用户数据告诉我：Life和Work场景的问题结构完全不同，混在一起模型会输出噪音。Decision场景需要对比逻辑，Reflect场景需要反思引导——这两个本质也不一样。"'),
]
for i, (label, color, _, desc) in enumerate(feedback_cards):
    y = Inches(2.3) + i * Inches(1.65)
    add_rounded_rect(slide, Inches(1), y, Inches(11), Inches(1.4), fill_color=BG_CARD, border_color=color)
    add_text_box(slide, Inches(1.3), y + Inches(0.15), Inches(3), Inches(0.3), label, font_size=15, color=color, bold=True)
    add_text_box(slide, Inches(1.3), y + Inches(0.55), Inches(10.2), Inches(0.7), desc, font_size=12, color=GRAY)

# ═══════════════ SLIDE 7: 语音 + 档案 ═══════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_section_label(slide, '06 · 特色功能')
add_title(slide, '语音通话 + 个人档案记忆')

# Left: Voice
add_rounded_rect(slide, Inches(1), Inches(2.5), Inches(5.3), Inches(4), fill_color=BG_CARD, border_color=BORDER)
add_text_box(slide, Inches(1.4), Inches(2.8), Inches(4.5), Inches(0.5), '🎙 语音通话模式', font_size=18, color=WHITE, bold=True)
voice_features = [
    '连续语音对话，像打电话一样',
    '3秒思考停顿自动发送，不打断思路',
    'AI回复TTS朗读，可以做家务时练',
    '支持Chrome / Safari（需HTTPS）',
    '文字模式 + 语音模式一键切换',
]
for i, f in enumerate(voice_features):
    add_text_box(slide, Inches(1.8), Inches(3.5) + i * Inches(0.45), Inches(4.2), Inches(0.35), f'• {f}', font_size=12, color=GRAY)

# Right: Profile
add_rounded_rect(slide, Inches(6.8), Inches(2.5), Inches(5.3), Inches(4), fill_color=BG_CARD, border_color=BORDER)
add_text_box(slide, Inches(7.2), Inches(2.8), Inches(4.5), Inches(0.5), '📝 个人档案记忆', font_size=18, color=WHITE, bold=True)
profile_features = [
    '背景信息注入AI system prompt',
    'Tempo / Mibo / Gap / 转行逻辑全记住',
    '每次练习不用重新交代背景',
    '可编辑档案面板，随时更新',
    '全部数据存本地localStorage',
]
for i, f in enumerate(profile_features):
    add_text_box(slide, Inches(7.6), Inches(3.5) + i * Inches(0.45), Inches(4.2), Inches(0.35), f'• {f}', font_size=12, color=GRAY)

# ═══════════════ SLIDE 8: 对比 ═══════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_section_label(slide, '07 · 为什么要做这个')
add_title(slide, '通用 AI 聊天 vs AI PM Coach')

headers = ['维度', '通用 AI 聊天', 'AI PM Coach']
rows = [
    ['记住你的项目', '需要每次重新交代', '✅ 档案持久注入 system prompt'],
    ['追问深度', '泛泛，不针对个人弱点', '✅ 知道你的薄弱点，定向深挖'],
    ['即时反馈', '不会主动点评每道题', '✅ 每答即评 ✅❌💡 三色标注'],
    ['语音通话', '不支持连续对话', '✅ 持续聆听 + TTS 朗读回复'],
    ['进度追踪', '无', '✅ 分数趋势 + 薄弱点标签'],
    ['面试模式', '需要手动构造 prompt', '✅ 创始人/HR/技术PM 一键切换'],
    ['数据隐私', '上传云端', '✅ 全部存本地 localStorage'],
]

# Table header
col_widths = [Inches(2.5), Inches(3.5), Inches(4.5)]
start_x = Inches(1.3)
header_y = Inches(2.5)
for j, (h, w) in enumerate(zip(headers, col_widths)):
    x = start_x + sum(col_widths[:j])
    add_rounded_rect(slide, x, header_y, w, Inches(0.5), fill_color=BRAND, border_color=None)
    add_text_box(slide, x + Inches(0.2), header_y + Inches(0.08), w - Inches(0.4), Inches(0.35), h, font_size=13, color=WHITE, bold=True)

# Table rows
for i, row in enumerate(rows):
    y = header_y + Inches(0.6) + i * Inches(0.55)
    bg = BG_CARD if i % 2 == 0 else BG_DARK
    for j, (cell, w) in enumerate(zip(row, col_widths)):
        x = start_x + sum(col_widths[:j])
        add_rounded_rect(slide, x, y, w, Inches(0.5), fill_color=bg, border_color=BORDER)
        c = WHITE if j == 0 else (GREEN if '✅' in cell else GRAY)
        add_text_box(slide, x + Inches(0.2), y + Inches(0.08), w - Inches(0.4), Inches(0.35), cell, font_size=12, color=c)

# ═══════════════ SLIDE 9: 技术架构 ═══════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_section_label(slide, '08 · 技术架构')
add_title(slide, '轻量、离线友好、手机优先')

tech_cards = [
    ('🧠', 'AI 引擎', 'DeepSeek Chat API\nSystem prompt 注入\n个人档案+面试官人设\n每轮明确指令'),
    ('🎤', '语音层', 'Web Speech API\nSpeechRecognition\n+ SpeechSynthesis\n3秒静默自动发送'),
    ('💾', '数据层', 'localStorage 全本地\n个人档案/API Key/\n面试记录/评分数据\n不上传任何服务器'),
    ('🚀', '部署', 'Vite 纯静态构建\nNetlify 自动部署\nGit push 即上线\n自带HTTPS免翻墙'),
]
for i, (icon, title, desc) in enumerate(tech_cards):
    x = Inches(1) + i * Inches(2.9)
    add_card(slide, x, Inches(2.5), Inches(2.6), Inches(3.6), icon, title, desc)

# ═══════════════ SLIDE 10: 路线图 ═══════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_section_label(slide, '09 · 产品路线图')
add_title(slide, '迭代方向')

roadmap = [
    ('P1', '实时语音打断', 'AI说话时可直接打断提问，更贴近真实面试节奏', AMBER),
    ('P2', '录音回放分析', '保存语音录音，AI分析语速、停顿频率、填充词使用量', AMBER),
    ('P3', '岗位匹配度评分', '粘贴JD后AI评估匹配度，生成针对性练习题目', AMBER),
    ('P4', '多人面试模拟', '创始人+HR+技术负责人同时在场，轮流提问', GRAY2),
]
for i, (level, title, desc, color) in enumerate(roadmap):
    y = Inches(2.5) + i * Inches(1.15)
    add_dot_line(slide, Inches(1), y + Inches(0.15), Inches(0.4), color)
    add_text_box(slide, Inches(1.6), y, Inches(1), Inches(0.3), level, font_size=13, color=color, bold=True)
    add_text_box(slide, Inches(2.5), y, Inches(3), Inches(0.3), title, font_size=16, color=WHITE, bold=True)
    add_text_box(slide, Inches(2.5), y + Inches(0.4), Inches(8), Inches(0.4), desc, font_size=12, color=GRAY)

# ═══════════════ SLIDE 11: Ending ═══════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_rect(slide, Inches(0), Inches(0), Inches(0.08), H, fill_color=BRAND)
add_text_box(slide, Inches(1), Inches(2.2), Inches(10), Inches(1),
    '用 AI 面试官，把每一次练习变成进步', font_size=36, color=WHITE, bold=True)
add_text_box(slide, Inches(1), Inches(3.4), Inches(10), Inches(0.5),
    'AI PM Coach v1.0  ·  刘美姣 独立设计与开发', font_size=16, color=GRAY)
add_text_box(slide, Inches(1), Inches(4.2), Inches(10), Inches(0.5),
    'React · DeepSeek API · Web Speech API · Netlify', font_size=14, color=GRAY2)
add_text_box(slide, Inches(1), Inches(5.5), Inches(10), Inches(0.5),
    'github.com/laumj0515-cloud/-stitch-ai-interview-muse', font_size=12, color=BRAND2)

# ── Save ──
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'AI_PM_Coach_产品说明书.pptx')
prs.save(output_path)
print(f'PPTX saved to: {output_path}')
